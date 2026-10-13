"""Gera o .docx do artigo a partir de docs/tcc/artigo/artigo.md.

    pip install python-docx pillow
    python scripts/artigo/gerar_docx.py   → docs/tcc/artigo/Artigo_PEI_Vivo_v2.docx

Formatação ABNT (NBR 14724/6022), na estrutura do modelo FAMETRO: capa, folha de
rosto, título com autores, resumo, abstract e seções numeradas. A4, margens
3/2 cm, Arial 12, entrelinha 1,5, recuo de 1,25 cm. Quadros, fontes, código e
referências em espaço simples. Código como listagem numerada (Código N);
diagramas com .svg ao lado do .png entram como vetor. Figuras com {paisagem} vão para página deitada; com {pagina}, para página retrato inteira.
Marcadores [PREENCHER: …] e [CONFERIR: …] saem destacados em amarelo.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAIZ = Path(__file__).resolve().parents[2]
ORIGEM = RAIZ / "docs/tcc/artigo/artigo.md"
DESTINO = RAIZ / "docs/tcc/artigo/Artigo_PEI_Vivo_v2.docx"

INSTITUICAO = ["CENTRO UNIVERSITÁRIO CEUNI – FAMETRO", "FACULDADE METROPOLITANA DE MANAUS – FAMETRO", "CURSO DE BACHARELADO EM SISTEMAS DE INFORMAÇÃO"]
AUTORES = ["GEANDRE ALFAIA COLARES", "JEAN VICTOR TORRES DOS SANTOS"]
NATUREZA = ("Artigo técnico apresentado ao Curso de Bacharelado em Sistemas de Informação da "
            "Faculdade Metropolitana de Manaus – FAMETRO, como requisito parcial para a obtenção "
            "do título de Bacharel em Sistemas de Informação.")
ORIENTADORA = "Orientadora: Profa. Luana Magalhães Leal"
LOCAL_ANO = ["MANAUS – AM", "2026"]
VINCULO = "Curso de Bacharelado em Sistemas de Informação, Centro Universitário CEUNI – FAMETRO, Manaus, Amazonas, Brasil"
AUTORIA = [  # nota de autoria (NBR 6022): função, curso, instituição, cidade, estado, país e e-mail
    ("Geandre Alfaia Colares", f"Discente do {VINCULO}. E-mail: geandre51@gmail.com"),
    ("Jean Victor Torres dos Santos", f"Discente do {VINCULO}. E-mail: jeanvictor.st29@gmail.com"),
    ("Luana Magalhães Leal", f"Docente e orientadora do {VINCULO}. E-mail: prof.luanalealm@gmail.com"),
]

FONTE = "Arial"
LARGURA_RETRATO = Cm(16)
MARCADOR = re.compile(r"(\[(?:PREENCHER|CONFERIR)[^\]]*\])")
NEGRITO = re.compile(r"(\*\*[^*]+\*\*)")


def espacamento(p, simples=False, antes=0, depois=0):
    f = p.paragraph_format
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE if simples else WD_LINE_SPACING.ONE_POINT_FIVE
    f.space_before, f.space_after = Pt(antes), Pt(depois)


def texto(p, s, tam=12, negrito=False, italico=False, mono=False):
    """Escreve s em p tratando **negrito** e destacando os marcadores de pendência."""
    for parte in MARCADOR.split(s):
        if not parte:
            continue
        pendente = bool(MARCADOR.fullmatch(parte))
        for trecho in ([parte] if pendente else NEGRITO.split(parte)):
            if not trecho:
                continue
            forte = trecho.startswith("**") and trecho.endswith("**")
            r = p.add_run(trecho[2:-2] if forte else trecho)
            r.font.name = "Consolas" if mono else FONTE
            r.font.size = Pt(tam)
            r.bold = negrito or forte
            r.italic = italico
            if pendente:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def paragrafo(doc, s, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=True, simples=False, tam=12, negrito=False, depois=0, antes=0):
    p = doc.add_paragraph()
    p.alignment = alinhar
    espacamento(p, simples, antes, depois)
    if recuo:
        p.paragraph_format.first_line_indent = Cm(1.25)
    return texto(p, s, tam=tam, negrito=negrito)


def centro(doc, s, negrito=True, tam=12, antes=0, depois=0):
    return paragrafo(doc, s, WD_ALIGN_PARAGRAPH.CENTER, recuo=False, negrito=negrito, tam=tam, antes=antes, depois=depois)


def quebra(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def numero_de_pagina(secao):
    p = secao.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run()
    for tipo, conteudo in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = conteudo
        r._r.append(el)
    r.font.name, r.font.size = FONTE, Pt(10)


def margens(secao, paisagem=False):
    secao.page_height, secao.page_width = (Cm(21), Cm(29.7)) if paisagem else (Cm(29.7), Cm(21))
    secao.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    secao.top_margin, secao.left_margin = Cm(3), Cm(3)
    secao.bottom_margin, secao.right_margin = Cm(2), Cm(2)


def capa_e_rosto(doc):
    for linha in INSTITUICAO:
        centro(doc, linha)
    for _ in range(4):
        centro(doc, "")
    for nome in AUTORES:
        centro(doc, nome)
    for _ in range(5):
        centro(doc, "")
    centro(doc, TITULO)
    for _ in range(9):
        centro(doc, "")
    for linha in LOCAL_ANO:
        centro(doc, linha)
    quebra(doc)

    for nome in AUTORES:
        centro(doc, nome)
    for _ in range(5):
        centro(doc, "")
    centro(doc, TITULO)
    for _ in range(3):
        centro(doc, "")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    espacamento(p, simples=True, depois=12)
    texto(p, NATUREZA, tam=10)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(8)
    espacamento(p, simples=True)
    texto(p, ORIENTADORA, tam=10)
    for _ in range(8):
        centro(doc, "")
    for linha in LOCAL_ANO:
        centro(doc, linha)
    quebra(doc)


def autores(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    espacamento(p, depois=6)
    for i, (nome, _) in enumerate(AUTORIA, 1):
        texto(p, nome)
        r = p.add_run(str(i))
        r.font.superscript, r.font.name, r.font.size = True, FONTE, Pt(12)
        if i < len(AUTORIA):
            p.add_run().add_break()
    for i, (_, nota) in enumerate(AUTORIA, 1):
        q = doc.add_paragraph()
        espacamento(q, simples=True)
        r = q.add_run(str(i))
        r.font.superscript, r.font.name, r.font.size = True, FONTE, Pt(10)
        texto(q, nota, tam=10)


def celula(c, s, cabecalho=False):
    c.paragraphs[0].text = ""
    p = c.paragraphs[0]
    espacamento(p, simples=True)
    texto(p, s, tam=10, negrito=cabecalho)


def tabela(doc, linhas):
    cab, *corpo = linhas
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, s in zip(t.rows[0].cells, cab):
        celula(c, s, cabecalho=True)
        sombra = OxmlElement("w:shd")
        sombra.set(qn("w:val"), "clear")
        sombra.set(qn("w:fill"), "E7EDF2")
        c._tc.get_or_add_tcPr().append(sombra)
    for linha in corpo:
        for c, s in zip(t.add_row().cells, linha):
            celula(c, s)
    # largura de cada coluna proporcional ao texto mais longo dela (mínimo 2,1 cm)
    pesos = [max(len(l[k]) for l in linhas) ** 0.8 for k in range(len(cab))]
    larguras = [max(Cm(2.1), int(LARGURA_RETRATO * w / sum(pesos))) for w in pesos]
    escala = LARGURA_RETRATO / sum(larguras)
    t.autofit = False
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), larguras):
        col.set(qn("w:w"), str(int(w * escala / 635)))  # EMU → twips
    for linha in t.rows:
        for c, w in zip(linha.cells, larguras):
            c.width = int(w * escala)
    # cabeçalho repete em nova página e nunca fica sozinho no fim dela
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for linha in t.rows[:2]:
        for c in linha.cells:
            for par in c.paragraphs:
                par.paragraph_format.keep_with_next = True


def legenda(doc, s):
    """Identificação de quadro ou figura, acima do elemento (ABNT)."""
    p = paragrafo(doc, s, WD_ALIGN_PARAGRAPH.CENTER, recuo=False, simples=True, tam=10, negrito=True, antes=6, depois=3)
    p.paragraph_format.keep_with_next = True


def fonte_do_elemento(doc, s):
    paragrafo(doc, s, WD_ALIGN_PARAGRAPH.CENTER, recuo=False, simples=True, tam=10, antes=3, depois=9)


def figura(doc, titulo, caminho, modo, fonte_texto):
    paisagem = modo == "paisagem"
    if modo == "pagina":  # diagrama alto: página retrato inteira
        quebra(doc)
        larg_max, alt_max = LARGURA_RETRATO, Cm(23)
    elif paisagem:
        nova = doc.add_section(WD_SECTION.NEW_PAGE)
        margens(nova, paisagem=True)
        larg_max, alt_max = Cm(24.7), Cm(13.5)
    else:
        larg_max, alt_max = LARGURA_RETRATO, Cm(10)
    legenda(doc, titulo)
    with Image.open(caminho) as img:
        w, h = img.size
    escala = min(larg_max / w, alt_max / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    espacamento(p, simples=True)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(caminho), width=int(w * escala))
    if caminho.with_suffix(".svg").exists():
        vetor(doc, r, caminho.with_suffix(".svg"))
    fonte_do_elemento(doc, fonte_texto)
    if paisagem:
        margens(doc.add_section(WD_SECTION.NEW_PAGE))


def vetor(doc, run, svg):
    """Embute o SVG junto do PNG (o Word 2016+ usa o vetor e o PNG fica de reserva):
    o diagrama pode ser ampliado no arquivo digital sem perder nitidez."""
    vetor.n = getattr(vetor, "n", 0) + 1
    parte = Part(PackURI(f"/word/media/vetor{vetor.n}.svg"), "image/svg+xml", svg.read_bytes(), doc.part.package)
    rid = doc.part.relate_to(parte, RT.IMAGE)
    ns_svg = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    lista, ext = OxmlElement("a:extLst"), OxmlElement("a:ext")
    ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = ext.makeelement(f"{{{ns_svg}}}svgBlip", nsmap={"asvg": ns_svg})
    svg_blip.set(qn("r:embed"), rid)
    ext.append(svg_blip)
    lista.append(ext)
    run._r.find(".//" + qn("a:blip")).append(lista)


def codigo(doc, linhas):
    """Listagem como texto (não imagem): fonte monoespaçada e linhas numeradas."""
    for i, linha in enumerate(linhas, 1):
        p = doc.add_paragraph()
        espacamento(p, simples=True, antes=3 if i == 1 else 0)
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.keep_with_next = i < len(linhas)
        r = p.add_run(f"{i:>2}  ")
        r.font.name, r.font.size = "Consolas", Pt(9)
        r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        texto(p, linha or " ", tam=9, mono=True)


def celulas(linha):
    return [c.strip() for c in linha.strip().strip("|").split("|")]


md = re.sub(r"<!--.*?-->", "", ORIGEM.read_text(encoding="utf-8"), flags=re.S)
linhas = md.splitlines()
TITULO = next(l[2:].strip() for l in linhas if l.startswith("# "))

doc = Document()
estilo = doc.styles["Normal"]
estilo.font.name, estilo.font.size = FONTE, Pt(12)
estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
margens(doc.sections[0])
numero_de_pagina(doc.sections[0])
capa_e_rosto(doc)

i, titulo_feito, referencias = 0, False, False
while i < len(linhas):
    l = linhas[i]
    if not l.strip():
        i += 1
        continue
    if l.startswith("# ") and not titulo_feito:
        centro(doc, TITULO, depois=12)
        titulo_feito = True
    elif l.strip() == "@autores":
        autores(doc)
    elif l.startswith("## ") and l[3:].strip() in ("RESUMO", "ABSTRACT"):
        if l[3:].strip() == "ABSTRACT":
            quebra(doc)
        centro(doc, l[3:].strip(), antes=12, depois=6)
    elif l.startswith("# "):
        referencias = l[2:].strip() == "REFERÊNCIAS"
        quebra(doc) if l[2:].startswith("1.") or referencias else None
        p = paragrafo(doc, l[2:].strip(), WD_ALIGN_PARAGRAPH.CENTER if referencias else WD_ALIGN_PARAGRAPH.LEFT, recuo=False, negrito=True, antes=12, depois=6)
        p.paragraph_format.keep_with_next = True
    elif l.startswith("## ") or l.startswith("### "):
        p = paragrafo(doc, l.lstrip("#").strip(), WD_ALIGN_PARAGRAPH.LEFT, recuo=False, negrito=l.startswith("## "), antes=6, depois=3)
        p.paragraph_format.keep_with_next = True
    elif re.match(r"^(Quadro|Figura|Código) \d+ – ", l):
        legenda(doc, l.strip())
    elif l.startswith("|"):
        bloco = []
        while i < len(linhas) and linhas[i].startswith("|"):
            if not re.fullmatch(r"\|[\s\-|:]+\|", linhas[i].strip()):
                bloco.append(celulas(linhas[i]))
            i += 1
        tabela(doc, bloco)
        continue
    elif l.startswith("!["):
        m = re.match(r"!\[(.+?)\]\((.+?)\)(?:\{(paisagem|pagina)\})?", l)
        j = i + 1
        while not linhas[j].strip():
            j += 1
        fonte_texto = linhas[j].strip() if linhas[j].startswith("Fonte:") else "Fonte: Autoria própria (2026)."
        figura(doc, m.group(1), (ORIGEM.parent / m.group(2)).resolve(), m.group(3), fonte_texto)
        i = j + 1 if linhas[j].startswith("Fonte:") else i + 1
        continue
    elif l.startswith("```"):
        j = i + 1
        while not linhas[j].startswith("```"):
            j += 1
        codigo(doc, linhas[i + 1 : j])
        i = j + 1
        continue
    elif l.startswith("Fonte:"):
        fonte_do_elemento(doc, l.strip())
    elif l.startswith("**Palavras-chave:**") or l.startswith("**Keywords:**"):
        paragrafo(doc, l, WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=False, simples=True, antes=6, depois=6)
    elif referencias:
        paragrafo(doc, l, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, simples=True, depois=12)
    elif re.match(r"^[a-e]\) ", l):
        paragrafo(doc, l, recuo=True)
    else:
        resumo = any(h.strip() in ("## RESUMO", "## ABSTRACT") for h in linhas[max(0, i - 2) : i])
        paragrafo(doc, l, recuo=not resumo, simples=resumo)
    i += 1

doc.save(DESTINO)
pendencias = len(MARCADOR.findall(md))
print(f"OK {DESTINO.relative_to(RAIZ)}  ({pendencias} marcadores [PREENCHER]/[CONFERIR] destacados)")
